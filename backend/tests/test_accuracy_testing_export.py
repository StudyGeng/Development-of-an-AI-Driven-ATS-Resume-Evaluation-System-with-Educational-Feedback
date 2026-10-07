from __future__ import annotations

import importlib
import os
import sys
import unittest
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import app_config  # noqa: E402


with patch.dict(
    os.environ,
    {
        "APP_ENV": "testing",
        "FLASK_SECRET_KEY": "accuracy-export-tests-only-secret-key-1234567890",
        "ACCURACY_TESTING_AUTO_EXPORT": "0",
    },
    clear=False,
):
    app_module = importlib.import_module("app")


class FakeCell:
    def __init__(self, value=None):
        self.value = value


class FakeSheet:
    def __init__(self, title: str = ""):
        self.title = title
        self.max_row = 0
        self.max_column = 0
        self._cells: dict[tuple[int, int], FakeCell] = {}

    def cell(self, row: int, column: int) -> FakeCell:
        self.max_row = max(self.max_row, row)
        self.max_column = max(self.max_column, column)
        return self._cells.setdefault((row, column), FakeCell())


class FakeWorkbook:
    def __init__(self, sheet: FakeSheet):
        self.sheet = sheet
        sheet.title = "Accuracy Testing"
        self._sheets = {"Accuracy Testing": sheet}

    @property
    def sheetnames(self) -> list[str]:
        return list(self._sheets)

    def __getitem__(self, name: str) -> FakeSheet:
        return self._sheets[name]

    def create_sheet(self, title: str) -> FakeSheet:
        if title in self._sheets:
            raise ValueError(f"Sheet {title!r} already exists")
        sheet = FakeSheet(title)
        self._sheets[title] = sheet
        return sheet


SIMPLIFIED_HEADERS = [
    "Test ID",
    "Resume Scenario",
    "Target Role",
    "Professional Resume Evidence",
    "Expected Answer",
    "ATS Score",
    "ATS Feedback",
    "System Answer",
    "Check",
    "Source File",
]

LEGACY_HEADERS = [
    "Resume ID",
    "Expected Role",
    "Recommended Role",
    "ATS Score",
    "Score OK?",
    "Keywords OK?",
    "Explanation OK?",
    "Keyword Notes",
    "Explanation Notes",
    "Role Check",
    "Keyword Check",
    "Explanation Check",
    "Status",
    "Discussion",
    "Source File",
]

COMBINED_HEADERS = [
    "Test ID",
    "Test Mode",
    "Selected Role",
    "Expected Role",
    "Evidence Basis",
    "Expected Fit",
    "ATS Score",
    "Recommended Role",
    "System Fit",
    "Check",
    "Source File",
]

HISTORY_HEADERS = [
    "Run #",
    "Tested At",
    "Test ID",
    "Test Mode",
    "Source File",
    "Selected Role",
    "Expected Role",
    "Evidence Basis",
    "Expected Fit",
    "ATS Score",
    "Recommended Role",
    "System Fit",
    "Result",
    "Recommendation Score",
    "Top 3 Role Matches",
    "Matched Keywords",
    "Missing Keywords",
    "Recommendation Reason",
]

CHECK_FORMULA = '=IF(OR(E7="",H7=""),"WAITING",IF(E7=H7,"PASS","REVIEW"))'
COMBINED_CHECK_FORMULA = (
    '=IF($A12="","",IF($B12="Selected role",'
    'IF(OR($C12="",$F12="",$I12=""),"WAITING",IF($F12=$I12,"PASS","REVIEW")),'
    'IF($B12="Auto-detect",IF(OR($D12="",$H12=""),"WAITING",'
    'IF(LOWER(TRIM($D12))=LOWER(TRIM($H12)),"PASS","REVIEW")),"SET MODE")))'
)


def set_row(sheet: FakeSheet, row: int, values: list[object]) -> None:
    for column, value in enumerate(values, start=1):
        sheet.cell(row, column).value = value


def add_history_sheet(
    workbook: Any,
    rows: list[list[object]] | None = None,
) -> Any:
    history = workbook.create_sheet("Test Run History")
    set_row(history, 1, HISTORY_HEADERS)
    for row_number, values in enumerate(rows or [], start=2):
        set_row(history, row_number, values)
    return history


def prepared_simplified_workbook(
    source_file: str = "ATS-01_Professional_Matching_Resume.pdf",
    *,
    test_id: str = "ATS-01",
    scenario: str = "Matching resume",
    evidence: str = "Professional Python, Flask, SQL, Git, testing, and debugging evidence.",
    expected_answer: str = "MATCH",
) -> tuple[FakeWorkbook, FakeSheet]:
    sheet = FakeSheet()
    set_row(sheet, 6, SIMPLIFIED_HEADERS)
    set_row(
        sheet,
        7,
        [
            test_id,
            scenario,
            "Software Developer",
            evidence,
            expected_answer,
            None,
            None,
            None,
            CHECK_FORMULA,
            source_file,
        ],
    )
    return FakeWorkbook(sheet), sheet


def prepared_legacy_workbook(source_file: str = "SD01.pdf") -> tuple[FakeWorkbook, FakeSheet]:
    sheet = FakeSheet()
    set_row(sheet, 4, LEGACY_HEADERS)
    set_row(
        sheet,
        5,
        [
            "SD01",
            "Software Developer",
            None,
            None,
            "No",
            "Yes",
            "No",
            None,
            None,
            None,
            None,
            None,
            None,
            "Human tester explanation",
            source_file,
        ],
    )
    return FakeWorkbook(sheet), sheet


def prepared_combined_workbook(
    source_file: str = "ATS-01_Professional_Matching_Resume.pdf",
    *,
    test_id: str = "SEL-01",
    test_mode: str = "Selected role",
    selected_role: str = "Software Developer",
    expected_role: str = "",
    evidence: str = "Professional Python, Flask, SQL, Git, testing, and debugging evidence.",
    expected_fit: str = "MATCH",
) -> tuple[Workbook, Any]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Accuracy Testing"
    set_row(sheet, 11, COMBINED_HEADERS)
    set_row(
        sheet,
        12,
        [
            test_id,
            test_mode,
            selected_role,
            expected_role,
            evidence,
            expected_fit,
            None,
            None,
            None,
            COMBINED_CHECK_FORMULA,
            source_file,
        ],
    )
    return workbook, sheet


def sample_analysis(
    uploaded_name: str = "ATS-01_Professional_Matching_Resume.pdf",
    *,
    matching_percentage: object = 82.5,
    recommended_role: str = "Software Engineer",
) -> dict:
    return {
        "uploaded_resume_name": uploaded_name,
        "matching_percentage": matching_percentage,
        "job_description_analysis": {
            "missing_keywords": ["testing"],
            "matched_keywords": ["python"],
        },
        "role_recommendation": {
            "recommended_role": recommended_role,
            "reason_summary": "The recommendation is supported by Python development evidence.",
        },
    }


def rich_selected_analysis(
    uploaded_name: str = "ATS-01_Professional_Matching_Resume.pdf",
    *,
    matching_percentage: object = 84.25,
    recommended_role: str = "Software Engineer",
) -> dict:
    analysis = sample_analysis(
        uploaded_name,
        matching_percentage=matching_percentage,
        recommended_role=recommended_role,
    )
    analysis.update(
        {
            "selected_role": "Software Developer",
            "target_role": "Software Developer",
            "top_role_matches": [
                {"role": recommended_role, "score": 91.5},
                {"role": "Full Stack Developer", "score": 82.25},
                {"role": "Cloud Engineer", "score": 70.0},
            ],
            "queued_at": "2026-09-29T10:11:12+08:00",
        }
    )
    analysis["job_description_analysis"] = {
        "missing_keywords": ["docker", "kubernetes"],
        "matched_keywords": ["python", "flask", "sql"],
    }
    analysis["role_recommendation"].update(
        {
            "selected_role": "Software Developer",
            "selected_database_role": "Software Developer",
            "recommended_database_role": recommended_role,
            "recommended_score": 91.5,
            "top_role_matches": analysis["top_role_matches"],
            "reason_summary": (
                "Strong Python, Flask, SQL, testing, and debugging evidence."
            ),
        }
    )
    return analysis


def unregistered_analysis(
    *,
    selected_role: str = "",
    matching_percentage: object = 80.96,
    recommended_role: str = "Computer Programmer",
    analysis_id: int | None = 189,
    tested_at: str = "2026-09-29T23:06:31+08:00",
) -> dict:
    analysis = rich_selected_analysis(
        "NGU MING HENN_Resume_FINAL.pdf",
        matching_percentage=matching_percentage,
        recommended_role=recommended_role,
    )
    analysis["analysis_id"] = analysis_id
    analysis["selected_role"] = selected_role
    analysis["target_role"] = selected_role or recommended_role
    analysis["accuracy_tested_at"] = tested_at
    analysis["role_recommendation"]["selected_role"] = selected_role
    analysis["role_recommendation"]["selected_database_role"] = selected_role
    return analysis


class AccuracyTestingConfigurationTests(unittest.TestCase):
    def test_export_defaults_off_in_development(self):
        with patch.dict(os.environ, {"APP_ENV": "development"}, clear=True):
            self.assertFalse(app_config.get_accuracy_testing_auto_export_enabled())

    def test_export_requires_explicit_local_opt_in(self):
        with patch.dict(
            os.environ,
            {"APP_ENV": "development", "ACCURACY_TESTING_AUTO_EXPORT": "1"},
            clear=True,
        ):
            self.assertTrue(app_config.get_accuracy_testing_auto_export_enabled())

    def test_export_never_enables_in_production(self):
        with patch.dict(
            os.environ,
            {"APP_ENV": "production", "ACCURACY_TESTING_AUTO_EXPORT": "1"},
            clear=True,
        ):
            self.assertFalse(app_config.get_accuracy_testing_auto_export_enabled())

    def test_invalid_export_setting_is_rejected(self):
        with patch.dict(
            os.environ,
            {"APP_ENV": "development", "ACCURACY_TESTING_AUTO_EXPORT": "maybe"},
            clear=True,
        ):
            with self.assertRaisesRegex(RuntimeError, "ACCURACY_TESTING_AUTO_EXPORT"):
                app_config.get_accuracy_testing_auto_export_enabled()


class AccuracyTestingWorkbookTests(unittest.TestCase):
    def test_combined_schema_is_discovered_from_row_eleven(self):
        _, sheet = prepared_combined_workbook()

        schema = app_module.get_accuracy_testing_schema(sheet)

        self.assertIsNotNone(schema)
        self.assertEqual(schema["kind"], "combined")
        self.assertEqual(schema["header_row"], 11)
        self.assertEqual(schema["data_end_row"], 31)
        self.assertEqual(
            schema["columns"],
            {
                "test_id": 1,
                "test_mode": 2,
                "selected_role": 3,
                "expected_role": 4,
                "evidence": 5,
                "expected_fit": 6,
                "ats_score": 7,
                "recommended_role": 8,
                "system_fit": 9,
                "check": 10,
                "source_file": 11,
            },
        )

    def test_combined_schema_uses_accuracy_cases_table_end_row(self):
        workbook, sheet = prepared_combined_workbook()
        sheet.add_table(Table(displayName="AccuracyCasesTable", ref="A11:K35"))

        schema = app_module.get_accuracy_testing_schema(sheet)

        self.assertIsNotNone(schema)
        self.assertEqual(schema["data_end_row"], 35)
        self.assertIs(workbook["Accuracy Testing"], sheet)

    def test_combined_case_beyond_legacy_row_limit_is_discoverable(self):
        workbook, sheet = prepared_combined_workbook()
        set_row(
            sheet,
            35,
            [
                "REV-AN-170",
                "Selected role",
                "Web Designer",
                "",
                "Independent human evidence.",
                "NOT MATCH",
                None,
                None,
                None,
                COMBINED_CHECK_FORMULA.replace("12", "35"),
                "software-developer_8.pdf",
            ],
        )
        sheet.add_table(Table(displayName="AccuracyCasesTable", ref="A11:K35"))
        analysis = sample_analysis(
            "software-developer_8.pdf",
            recommended_role="Full Stack Developer",
        )
        analysis["selected_role"] = "Web Designer"

        found_sheet, row, warning = app_module.find_accuracy_testing_row(
            workbook,
            analysis,
        )

        self.assertIs(found_sheet, sheet)
        self.assertEqual((row, warning), (35, ""))

    def test_combined_validation_is_mode_specific(self):
        workbook, sheet = prepared_combined_workbook()
        schema = app_module.get_accuracy_testing_schema(sheet)

        self.assertEqual(app_module.validate_accuracy_testing_case(sheet, schema, 12), "")

        sheet.cell(12, 3).value = ""
        self.assertIn(
            "Selected Role",
            app_module.validate_accuracy_testing_case(sheet, schema, 12),
        )
        sheet.cell(12, 3).value = "Software Developer"
        sheet.cell(12, 6).value = "MAYBE"
        self.assertIn(
            "MATCH or NOT MATCH",
            app_module.validate_accuracy_testing_case(sheet, schema, 12),
        )
        sheet.cell(12, 6).value = "MATCH"
        sheet.cell(12, 5).value = ""
        self.assertIn(
            "Evidence Basis",
            app_module.validate_accuracy_testing_case(sheet, schema, 12),
        )

        sheet.cell(12, 2).value = "Auto-detect"
        sheet.cell(12, 3).value = ""
        sheet.cell(12, 4).value = "Software Developer"
        sheet.cell(12, 5).value = "Independent human evidence."
        sheet.cell(12, 6).value = ""
        self.assertEqual(app_module.validate_accuracy_testing_case(sheet, schema, 12), "")
        sheet.cell(12, 5).value = ""
        self.assertIn(
            "Evidence Basis",
            app_module.validate_accuracy_testing_case(sheet, schema, 12),
        )

        sheet.cell(12, 4).value = ""
        self.assertIn(
            "Expected Role",
            app_module.validate_accuracy_testing_case(sheet, schema, 12),
        )
        sheet.cell(12, 4).value = "Software Developer"
        sheet.cell(12, 3).value = "Software Developer"
        self.assertIn(
            "must leave Selected Role and Expected Fit blank",
            app_module.validate_accuracy_testing_case(sheet, schema, 12),
        )

        self.assertIs(workbook["Accuracy Testing"], sheet)

    def test_same_source_file_routes_by_mode_and_selected_role(self):
        shared_source = "professional-software-resume.pdf"
        workbook, sheet = prepared_combined_workbook(source_file=shared_source)
        set_row(
            sheet,
            13,
            [
                "SEL-02",
                "Selected role",
                "Cybersecurity Analyst",
                "",
                "Security comparison case.",
                "NOT MATCH",
                None,
                None,
                None,
                COMBINED_CHECK_FORMULA.replace("12", "13"),
                shared_source,
            ],
        )
        set_row(
            sheet,
            14,
            [
                "AUTO-01",
                "Auto-detect",
                "",
                "Software Developer",
                "Exact top-one role case.",
                "",
                None,
                None,
                None,
                COMBINED_CHECK_FORMULA.replace("12", "14"),
                shared_source,
            ],
        )

        selected_software = sample_analysis(shared_source)
        selected_software["selected_role"] = "Software Developer"
        _, software_row, software_warning = app_module.find_accuracy_testing_row(
            workbook, selected_software
        )

        selected_security = sample_analysis(shared_source)
        selected_security["selected_role"] = "Cybersecurity Analyst"
        _, security_row, security_warning = app_module.find_accuracy_testing_row(
            workbook, selected_security
        )

        auto_detect = sample_analysis(shared_source)
        auto_detect["selected_role"] = ""
        _, auto_row, auto_warning = app_module.find_accuracy_testing_row(
            workbook, auto_detect
        )

        self.assertEqual((software_row, software_warning), (12, ""))
        self.assertEqual((security_row, security_warning), (13, ""))
        self.assertEqual((auto_row, auto_warning), (14, ""))

    def test_combined_duplicate_source_mode_and_selected_role_is_rejected(self):
        duplicate_name = "duplicate-composite.pdf"
        workbook, sheet = prepared_combined_workbook(source_file=duplicate_name)
        set_row(
            sheet,
            13,
            [
                "SEL-02",
                "Selected role",
                "Software Developer",
                "",
                "A second prepared case with the same composite key.",
                "MATCH",
                None,
                None,
                None,
                COMBINED_CHECK_FORMULA.replace("12", "13"),
                duplicate_name,
            ],
        )
        analysis = sample_analysis(duplicate_name)
        analysis["selected_role"] = "Software Developer"

        found_sheet, row, warning = app_module.find_accuracy_testing_row(
            workbook, analysis
        )

        self.assertIs(found_sheet, sheet)
        self.assertIsNone(row)
        self.assertIn("same source, mode, and selected role", warning)

        result, write_warning = app_module.write_accuracy_testing_analysis(
            workbook,
            analysis,
        )

        self.assertIsNone(result)
        self.assertIn("same source, mode, and selected role", write_warning)
        self.assertNotIn("Test Run History", workbook.sheetnames)

    def test_combined_selected_write_populates_g_h_i_and_preserves_prepared_cells(self):
        workbook, sheet = prepared_combined_workbook()
        preserved_columns = (1, 2, 3, 4, 5, 6, 10, 11)
        prepared_values = {
            column: sheet.cell(12, column).value for column in preserved_columns
        }
        analysis = sample_analysis(
            matching_percentage=84.25,
            recommended_role="Software Engineer",
        )
        analysis["selected_role"] = "Software Developer"

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            analysis,
        )

        self.assertEqual((row, warning), (12, ""))
        self.assertEqual(sheet.cell(12, 7).value, 84.25)
        self.assertEqual(sheet.cell(12, 8).value, "Software Engineer")
        self.assertEqual(sheet.cell(12, 9).value, "MATCH")
        for column, value in prepared_values.items():
            self.assertEqual(sheet.cell(12, column).value, value)

    def test_selected_run_history_records_all_eighteen_fields(self):
        workbook, sheet = prepared_combined_workbook()
        analysis = rich_selected_analysis()

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            analysis,
        )

        self.assertEqual((row, warning), (12, ""))
        history = workbook["Test Run History"]
        self.assertEqual(
            [history.cell(1, column).value for column in range(1, 19)],
            HISTORY_HEADERS,
        )
        values = [history.cell(2, column).value for column in range(1, 19)]
        self.assertEqual(values[0], 1)
        self.assertIsInstance(values[1], datetime)
        self.assertEqual(values[2], "SEL-01")
        self.assertEqual(values[3], "Selected role")
        self.assertEqual(values[4], "ATS-01_Professional_Matching_Resume.pdf")
        self.assertEqual(values[5], "Software Developer")
        self.assertEqual(values[6], "")
        self.assertEqual(
            values[7],
            "Professional Python, Flask, SQL, Git, testing, and debugging evidence.",
        )
        self.assertEqual(values[8], "MATCH")
        self.assertEqual(values[9], 84.25)
        self.assertEqual(values[10], "Software Engineer")
        self.assertEqual(values[11], "MATCH")
        self.assertEqual(values[12], "PASS")
        self.assertEqual(values[13], 91.5)
        self.assertEqual(
            values[14],
            (
                "1. Software Engineer (91.50); "
                "2. Full Stack Developer (82.25); "
                "3. Cloud Engineer (70.00)"
            ),
        )
        self.assertEqual(values[15], "python, flask, sql")
        self.assertEqual(values[16], "docker, kubernetes")
        self.assertEqual(
            values[17],
            "Strong Python, Flask, SQL, testing, and debugging evidence.",
        )
        self.assertEqual(history.tables["TestRunHistoryTable"].ref, "A1:R2")

    def test_repeated_combined_runs_append_history_and_main_row_keeps_latest(self):
        workbook, sheet = prepared_combined_workbook()
        first_analysis = rich_selected_analysis(
            matching_percentage=81.0,
            recommended_role="Software Developer",
        )
        second_analysis = rich_selected_analysis(
            matching_percentage=93.75,
            recommended_role="Full Stack Developer",
        )

        first_result = app_module.write_accuracy_testing_analysis_row(
            workbook,
            first_analysis,
        )
        second_result = app_module.write_accuracy_testing_analysis_row(
            workbook,
            second_analysis,
        )

        self.assertEqual(first_result, (12, ""))
        self.assertEqual(second_result, (12, ""))
        self.assertEqual(sheet.cell(12, 7).value, 93.75)
        self.assertEqual(sheet.cell(12, 8).value, "Full Stack Developer")
        self.assertEqual(sheet.cell(12, 9).value, "MATCH")
        history = workbook["Test Run History"]
        self.assertEqual(history.cell(2, 1).value, 1)
        self.assertEqual(history.cell(3, 1).value, 2)
        self.assertEqual(history.cell(2, 10).value, 81.0)
        self.assertEqual(history.cell(3, 10).value, 93.75)
        self.assertEqual(history.cell(2, 11).value, "Software Developer")
        self.assertEqual(history.cell(3, 11).value, "Full Stack Developer")
        self.assertEqual(history.tables["TestRunHistoryTable"].ref, "A1:R3")

    def test_unregistered_ngu_auto_analysis_appends_unreviewed_history_only(self):
        workbook, sheet = prepared_combined_workbook()
        prepared_values = [sheet.cell(12, column).value for column in range(1, 12)]
        analysis = unregistered_analysis()

        result, warning = app_module.write_accuracy_testing_analysis(
            workbook,
            analysis,
        )

        self.assertEqual(warning, "")
        self.assertEqual(
            result,
            {
                "prepared_row": None,
                "history_row": 2,
                "history_only": True,
                "history_result": "UNREVIEWED",
            },
        )
        self.assertEqual(
            [sheet.cell(12, column).value for column in range(1, 12)],
            prepared_values,
        )
        history = workbook["Test Run History"]
        values = [history.cell(2, column).value for column in range(1, 19)]
        self.assertEqual(values[0], 1)
        self.assertEqual(values[1], datetime(2026, 9, 29, 23, 6, 31))
        self.assertEqual(values[2], "AN-189")
        self.assertEqual(values[3], "Auto-detect")
        self.assertEqual(values[4], "NGU MING HENN_Resume_FINAL.pdf")
        self.assertEqual(values[5:9], ["", "", "", ""])
        self.assertEqual(values[9], 80.96)
        self.assertEqual(values[10], "Computer Programmer")
        self.assertEqual(values[11], "")
        self.assertEqual(values[12], "UNREVIEWED")
        self.assertEqual(values[13], 91.5)
        self.assertIn("Computer Programmer (91.50)", values[14])
        self.assertEqual(values[15], "python, flask, sql")
        self.assertEqual(values[16], "docker, kubernetes")
        self.assertEqual(
            values[17],
            "Strong Python, Flask, SQL, testing, and debugging evidence.",
        )
        self.assertEqual(history.cell(2, 13).fill.fgColor.rgb[-6:], "FFF2CC")
        self.assertEqual(history.cell(2, 13).font.color.rgb[-6:], "9C6500")
        self.assertEqual(history.tables["TestRunHistoryTable"].ref, "A1:R2")

    def test_unregistered_selected_analysis_records_system_fit_without_ground_truth(self):
        workbook, sheet = prepared_combined_workbook()
        prepared_values = [sheet.cell(12, column).value for column in range(1, 12)]
        analysis = unregistered_analysis(
            selected_role="Software Developer",
            recommended_role="Software Engineer",
            analysis_id=None,
        )

        first_result, first_warning = app_module.write_accuracy_testing_analysis(
            workbook,
            analysis,
        )
        second_result, second_warning = app_module.write_accuracy_testing_analysis(
            workbook,
            analysis,
        )

        self.assertEqual(first_warning, "")
        self.assertEqual(second_warning, "")
        self.assertTrue(first_result["history_only"])
        self.assertTrue(second_result["history_only"])
        self.assertEqual(first_result["history_row"], 2)
        self.assertEqual(second_result["history_row"], 3)
        self.assertEqual(
            [sheet.cell(12, column).value for column in range(1, 12)],
            prepared_values,
        )
        history = workbook["Test Run History"]
        for history_row in (2, 3):
            self.assertEqual(history.cell(history_row, 3).value, "")
            self.assertEqual(history.cell(history_row, 4).value, "Selected role")
            self.assertEqual(history.cell(history_row, 5).value, "NGU MING HENN_Resume_FINAL.pdf")
            self.assertEqual(history.cell(history_row, 6).value, "Software Developer")
            self.assertEqual(history.cell(history_row, 7).value, "")
            self.assertEqual(history.cell(history_row, 8).value, "")
            self.assertEqual(history.cell(history_row, 9).value, "")
            self.assertEqual(history.cell(history_row, 11).value, "Software Engineer")
            self.assertEqual(history.cell(history_row, 12).value, "MATCH")
            self.assertEqual(history.cell(history_row, 13).value, "UNREVIEWED")
        self.assertEqual(history.tables["TestRunHistoryTable"].ref, "A1:R3")

    def test_unregistered_combined_invalid_score_appends_nothing(self):
        workbook, sheet = prepared_combined_workbook()
        prepared_values = [sheet.cell(12, column).value for column in range(1, 12)]
        analysis = unregistered_analysis(matching_percentage="not-a-score")

        result, warning = app_module.write_accuracy_testing_analysis(
            workbook,
            analysis,
        )

        self.assertIsNone(result)
        self.assertIn("valid ATS score", warning)
        self.assertEqual(
            [sheet.cell(12, column).value for column in range(1, 12)],
            prepared_values,
        )
        self.assertNotIn("Test Run History", workbook.sheetnames)

    def test_auto_history_result_requires_exact_top_role(self):
        workbook, sheet = prepared_combined_workbook(
            test_id="AUTO-01",
            test_mode="Auto-detect",
            selected_role="",
            expected_role="Software Developer",
            expected_fit="",
        )
        exact_analysis = sample_analysis(
            matching_percentage=88.0,
            recommended_role="  software   developer  ",
        )
        exact_analysis["selected_role"] = ""
        near_match_analysis = sample_analysis(
            matching_percentage=86.0,
            recommended_role="Software Engineer",
        )
        near_match_analysis["selected_role"] = ""

        exact_result = app_module.write_accuracy_testing_analysis_row(
            workbook,
            exact_analysis,
        )
        near_match_result = app_module.write_accuracy_testing_analysis_row(
            workbook,
            near_match_analysis,
        )

        self.assertEqual(exact_result, (12, ""))
        self.assertEqual(near_match_result, (12, ""))
        self.assertEqual(sheet.cell(12, 7).value, 86.0)
        self.assertEqual(sheet.cell(12, 8).value, "Software Engineer")
        self.assertIsNone(sheet.cell(12, 9).value)
        history = workbook["Test Run History"]
        self.assertEqual(history.cell(2, 4).value, "Auto-detect")
        self.assertEqual(history.cell(2, 13).value, "PASS")
        self.assertEqual(history.cell(3, 13).value, "REVIEW")

    def test_history_position_uses_max_numeric_run_and_last_nonblank_run_row(self):
        workbook, _ = prepared_combined_workbook()
        history = add_history_sheet(workbook)
        history.cell(2, 1).value = 4
        history.cell(3, 1).value = "manual-note"
        history.cell(5, 1).value = 9
        history.cell(6, 2).value = "content outside the append key"

        next_run, next_row = app_module.get_next_accuracy_history_position(history)

        self.assertEqual(next_run, 10)
        self.assertEqual(next_row, 6)

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            rich_selected_analysis(),
        )

        self.assertEqual((row, warning), (12, ""))
        self.assertEqual(history.cell(6, 1).value, 10)
        self.assertIsInstance(history.cell(6, 2).value, datetime)
        self.assertEqual(history.cell(6, 3).value, "SEL-01")

    def test_combined_auto_write_populates_g_h_and_clears_i(self):
        workbook, sheet = prepared_combined_workbook(
            test_id="AUTO-01",
            test_mode="Auto-detect",
            selected_role="",
            expected_role="Software Developer",
            expected_fit="",
        )
        sheet.cell(12, 9).value = "STALE VALUE"
        preserved_columns = (1, 2, 3, 4, 5, 6, 10, 11)
        prepared_values = {
            column: sheet.cell(12, column).value for column in preserved_columns
        }
        analysis = sample_analysis(
            matching_percentage=79,
            recommended_role="Software Developer",
        )
        analysis["selected_role"] = ""

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            analysis,
        )

        self.assertEqual((row, warning), (12, ""))
        self.assertEqual(sheet.cell(12, 7).value, 79.0)
        self.assertEqual(sheet.cell(12, 8).value, "Software Developer")
        self.assertIsNone(sheet.cell(12, 9).value)
        for column, value in prepared_values.items():
            self.assertEqual(sheet.cell(12, column).value, value)

    def test_combined_invalid_score_causes_no_partial_write(self):
        workbook, sheet = prepared_combined_workbook()
        history = add_history_sheet(workbook)
        sheet.cell(12, 7).value = 71.0
        sheet.cell(12, 8).value = "Existing Recommendation"
        sheet.cell(12, 9).value = "NOT MATCH"
        analysis = sample_analysis(
            matching_percentage="invalid score",
            recommended_role="Software Engineer",
        )
        analysis["selected_role"] = "Software Developer"

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            analysis,
        )

        self.assertIsNone(row)
        self.assertIn("valid ATS score", warning)
        self.assertEqual(sheet.cell(12, 7).value, 71.0)
        self.assertEqual(sheet.cell(12, 8).value, "Existing Recommendation")
        self.assertEqual(sheet.cell(12, 9).value, "NOT MATCH")
        self.assertEqual(sheet.cell(12, 10).value, COMBINED_CHECK_FORMULA)
        self.assertEqual(history.max_row, 1)

    def test_compact_queued_payload_retains_mode_and_role_fields(self):
        analysis = sample_analysis(recommended_role="Machine Learning Engineer")
        analysis.update(
            {
                "analysis_id": 321,
                "target_role": "Machine Learning Engineer",
                "selected_role": "Data Scientist",
                "queued_at": "2026-09-29T10:11:12+08:00",
                "top_role_matches": [
                    {"role": "Machine Learning Engineer", "score": 91.5},
                    {"role": "Data Scientist", "score": 84.0},
                ],
            }
        )
        analysis["role_recommendation"].update(
            {
                "selected_role": "Data Scientist",
                "selected_database_role": "Data Scientist",
                "recommended_database_role": "Machine Learning Engineer",
                "recommended_score": 91.5,
            }
        )

        compact = app_module.compact_accuracy_testing_analysis(analysis)

        self.assertEqual(compact["analysis_id"], 321)
        self.assertEqual(compact["accuracy_tested_at"], analysis["queued_at"])
        self.assertEqual(compact["target_role"], "Machine Learning Engineer")
        self.assertEqual(compact["selected_role"], "Data Scientist")
        self.assertEqual(
            compact["role_recommendation"]["selected_role"],
            "Data Scientist",
        )
        self.assertEqual(
            compact["role_recommendation"]["selected_database_role"],
            "Data Scientist",
        )
        self.assertEqual(
            compact["role_recommendation"]["recommended_role"],
            "Machine Learning Engineer",
        )
        self.assertEqual(
            compact["role_recommendation"]["recommended_database_role"],
            "Machine Learning Engineer",
        )
        self.assertEqual(compact["role_recommendation"]["recommended_score"], 91.5)
        self.assertEqual(compact["top_role_matches"], analysis["top_role_matches"])

    def test_compact_queued_payload_is_sufficient_for_later_history_row(self):
        workbook, _ = prepared_combined_workbook()
        compact = app_module.compact_accuracy_testing_analysis(
            rich_selected_analysis()
        )

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            compact,
        )

        self.assertEqual((row, warning), (12, ""))
        history = workbook["Test Run History"]
        self.assertEqual(history.cell(2, 3).value, "SEL-01")
        self.assertEqual(history.cell(2, 4).value, "Selected role")
        self.assertEqual(history.cell(2, 10).value, 84.25)
        self.assertEqual(history.cell(2, 11).value, "Software Engineer")
        self.assertEqual(history.cell(2, 13).value, "PASS")
        self.assertEqual(history.cell(2, 14).value, 91.5)
        self.assertIn("Software Engineer (91.50)", history.cell(2, 15).value)
        self.assertEqual(history.cell(2, 16).value, "python, flask, sql")
        self.assertEqual(history.cell(2, 17).value, "docker, kubernetes")
        self.assertEqual(
            history.cell(2, 18).value,
            "Strong Python, Flask, SQL, testing, and debugging evidence.",
        )

    def test_simplified_schema_is_discovered_from_row_six_with_source_file_in_j(self):
        _, sheet = prepared_simplified_workbook()

        schema = app_module.get_accuracy_testing_schema(sheet)

        self.assertIsNotNone(schema)
        self.assertEqual(schema["kind"], "simplified")
        self.assertEqual(schema["header_row"], 6)
        self.assertEqual(schema["columns"]["ats_score"], 6)
        self.assertEqual(schema["columns"]["ats_feedback"], 7)
        self.assertEqual(schema["columns"]["system_answer"], 8)
        self.assertEqual(schema["columns"]["check"], 9)
        self.assertEqual(schema["columns"]["source_file"], 10)

    def test_simplified_write_populates_system_outputs_and_preserves_prepared_cells(self):
        workbook, sheet = prepared_simplified_workbook()
        preserved_columns = (1, 2, 3, 4, 5, 9, 10)
        prepared_values = {
            column: sheet.cell(7, column).value for column in preserved_columns
        }

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            sample_analysis(),
        )

        self.assertEqual((row, warning), (7, ""))
        self.assertEqual(sheet.cell(7, 6).value, 82.5)
        self.assertIsInstance(sheet.cell(7, 6).value, float)
        self.assertIn("Recommended role: Software Engineer", sheet.cell(7, 7).value)
        self.assertIn("Python development evidence", sheet.cell(7, 7).value)
        self.assertIn("Missing: testing", sheet.cell(7, 7).value)
        self.assertIn("Matched: python", sheet.cell(7, 7).value)
        self.assertEqual(sheet.cell(7, 8).value, "MATCH")
        for column, value in prepared_values.items():
            self.assertEqual(sheet.cell(7, column).value, value)
        self.assertNotIn("Test Run History", workbook.sheetnames)

    def test_simplified_non_matching_role_writes_not_match(self):
        workbook, sheet = prepared_simplified_workbook(
            "ATS-02_Professional_Non_Matching_Resume.pdf",
            test_id="ATS-02",
            scenario="Non-matching resume",
            evidence="Professional business administration and retail operations evidence.",
            expected_answer="NOT MATCH",
        )

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            sample_analysis(
                "ATS-02_Professional_Non_Matching_Resume.pdf",
                matching_percentage=23,
                recommended_role="Business Data Analyst",
            ),
        )

        self.assertEqual((row, warning), (7, ""))
        self.assertEqual(sheet.cell(7, 6).value, 23.0)
        self.assertIn("Recommended role: Business Data Analyst", sheet.cell(7, 7).value)
        self.assertEqual(sheet.cell(7, 8).value, "NOT MATCH")
        self.assertEqual(sheet.cell(7, 5).value, "NOT MATCH")
        self.assertEqual(sheet.cell(7, 9).value, CHECK_FORMULA)

    def test_source_file_selects_only_the_prepared_simplified_case(self):
        workbook, _ = prepared_simplified_workbook("professional-software-resume.pdf")

        sheet, row, warning = app_module.find_accuracy_testing_row(
            workbook,
            sample_analysis("professional-software-resume.pdf"),
        )

        self.assertIsNotNone(sheet)
        self.assertEqual((row, warning), (7, ""))

    def test_filename_id_inference_supports_ats_01_variants(self):
        workbook, _ = prepared_simplified_workbook("different-file-name.pdf")

        sheet, row, warning = app_module.find_accuracy_testing_row(
            workbook,
            sample_analysis("ATS_01_candidate_resume.pdf"),
        )

        self.assertIsNotNone(sheet)
        self.assertEqual((row, warning), (7, ""))

    def test_duplicate_source_file_is_rejected(self):
        duplicate_name = "duplicate-professional-resume.pdf"
        workbook, sheet = prepared_simplified_workbook(duplicate_name)
        set_row(
            sheet,
            8,
            [
                "ATS-02",
                "Non-matching resume",
                "Software Developer",
                "Professional business and retail operations evidence.",
                "NOT MATCH",
                None,
                None,
                None,
                '=IF(OR(E8="",H8=""),"WAITING",IF(E8=H8,"PASS","REVIEW"))',
                duplicate_name,
            ],
        )

        found_sheet, row, warning = app_module.find_accuracy_testing_row(
            workbook,
            sample_analysis(duplicate_name),
        )

        self.assertIs(found_sheet, sheet)
        self.assertIsNone(row)
        self.assertIn("more than one prepared test case", warning)

    def test_unregistered_resume_does_not_create_a_test_case(self):
        workbook, sheet = prepared_simplified_workbook()

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            sample_analysis("unregistered-resume.pdf"),
        )

        self.assertIsNone(row)
        self.assertIn("not registered", warning)
        self.assertEqual(sheet.cell(7, 1).value, "ATS-01")
        for column in (6, 7, 8):
            self.assertIsNone(sheet.cell(7, column).value)
        self.assertIsNone(sheet.cell(8, 1).value)

    def test_missing_or_invalid_score_does_not_write_system_outputs(self):
        invalid_scores = {
            "missing": None,
            "non-numeric": "not-a-score",
            "nan": float("nan"),
            "infinite": float("inf"),
        }
        for label, invalid_score in invalid_scores.items():
            with self.subTest(label=label):
                workbook, sheet = prepared_simplified_workbook()
                analysis = sample_analysis(matching_percentage=invalid_score)
                if label == "missing":
                    analysis.pop("matching_percentage")

                row, warning = app_module.write_accuracy_testing_analysis_row(
                    workbook,
                    analysis,
                )

                self.assertIsNone(row)
                self.assertIn("valid ATS score", warning)
                for column in (6, 7, 8):
                    self.assertIsNone(sheet.cell(7, column).value)
                self.assertEqual(sheet.cell(7, 9).value, CHECK_FORMULA)

    def test_legacy_schema_remains_compatible_and_preserves_human_review(self):
        workbook, sheet = prepared_legacy_workbook()
        expected_review = {
            5: "No",
            6: "Yes",
            7: "No",
            14: "Human tester explanation",
        }

        row, warning = app_module.write_accuracy_testing_analysis_row(
            workbook,
            sample_analysis(
                "SD01.pdf",
                matching_percentage=82.5,
                recommended_role="Software Developer",
            ),
        )

        self.assertEqual((row, warning), (5, ""))
        self.assertEqual(sheet.cell(5, 3).value, "Software Developer")
        self.assertEqual(sheet.cell(5, 4).value, 82.5)
        self.assertIn("Missing: testing", sheet.cell(5, 8).value)
        self.assertIn("Python development evidence", sheet.cell(5, 9).value)
        for column, value in expected_review.items():
            self.assertEqual(sheet.cell(5, column).value, value)
        self.assertNotIn("Test Run History", workbook.sheetnames)

    def test_unregistered_real_file_lock_queues_then_flushes_history_only_runs(self):
        with TemporaryDirectory() as temporary_directory:
            workbook_path = Path(temporary_directory) / "ats_resume_fyp_testing.xlsx"
            pending_path = Path(temporary_directory) / "accuracy_testing_pending.jsonl"
            workbook, _ = prepared_combined_workbook()
            workbook.save(workbook_path)
            workbook.close()

            queued_analysis = unregistered_analysis(
                analysis_id=189,
                tested_at="2026-09-29T20:11:05+08:00",
            )
            current_analysis = unregistered_analysis(
                analysis_id=190,
                tested_at="2026-09-29T23:06:31+08:00",
            )
            configuration_patches = (
                patch.object(app_module, "ACCURACY_TESTING_WORKBOOK", str(workbook_path)),
                patch.object(app_module, "ACCURACY_TESTING_PENDING_FILE", str(pending_path)),
                patch.object(app_module, "is_production_environment", return_value=False),
                patch.object(app_module, "get_accuracy_testing_auto_export_enabled", return_value=True),
            )

            with configuration_patches[0], configuration_patches[1], configuration_patches[2], configuration_patches[3]:
                with patch(
                    "openpyxl.workbook.workbook.Workbook.save",
                    side_effect=PermissionError("workbook is locked"),
                ):
                    queued_result = app_module.update_accuracy_testing_workbook(
                        queued_analysis
                    )

                self.assertFalse(queued_result["accuracy_testing_saved"])
                self.assertTrue(queued_result["accuracy_testing_attempted"])
                self.assertTrue(queued_result["accuracy_testing_pending"])
                self.assertTrue(queued_result["accuracy_testing_history_only"])
                self.assertTrue(pending_path.exists())

                saved_result = app_module.update_accuracy_testing_workbook(
                    current_analysis
                )

            self.assertTrue(saved_result["accuracy_testing_saved"])
            self.assertTrue(saved_result["accuracy_testing_attempted"])
            self.assertIsNone(saved_result["accuracy_testing_row"])
            self.assertEqual(saved_result["accuracy_testing_history_row"], 3)
            self.assertTrue(saved_result["accuracy_testing_history_only"])
            self.assertEqual(
                saved_result["accuracy_testing_history_result"],
                "UNREVIEWED",
            )
            self.assertEqual(saved_result["accuracy_testing_saved_rows"], [])
            self.assertEqual(
                saved_result["accuracy_testing_saved_history_rows"],
                [2, 3],
            )
            self.assertEqual(saved_result["accuracy_testing_pending_flushed"], 1)
            self.assertFalse(pending_path.exists())

            saved_workbook = load_workbook(workbook_path)
            try:
                prepared_sheet = saved_workbook["Accuracy Testing"]
                for column in (7, 8, 9):
                    self.assertIsNone(prepared_sheet.cell(12, column).value)
                history = saved_workbook["Test Run History"]
                self.assertEqual(history.max_row, 3)
                self.assertEqual(history.cell(2, 2).value, datetime(2026, 9, 29, 20, 11, 5))
                self.assertEqual(history.cell(3, 2).value, datetime(2026, 9, 29, 23, 6, 31))
                self.assertEqual(history.cell(2, 3).value, "AN-189")
                self.assertEqual(history.cell(3, 3).value, "AN-190")
                self.assertEqual(history.cell(2, 13).value, "UNREVIEWED")
                self.assertEqual(history.cell(3, 13).value, "UNREVIEWED")
                self.assertEqual(
                    history.tables["TestRunHistoryTable"].ref,
                    "A1:R3",
                )
            finally:
                saved_workbook.close()

    def test_production_export_stops_before_filesystem_or_queue_access(self):
        with patch.dict(
            os.environ,
            {"APP_ENV": "production", "ACCURACY_TESTING_AUTO_EXPORT": "1"},
            clear=False,
        ), patch.object(app_module.os.path, "exists") as path_exists, patch.object(
            app_module, "load_pending_accuracy_testing_analyses"
        ) as load_pending, patch.object(
            app_module, "queue_accuracy_testing_analysis"
        ) as queue_analysis:
            result = app_module.update_accuracy_testing_workbook(sample_analysis())

        self.assertFalse(result["accuracy_testing_saved"])
        self.assertNotIn("accuracy_testing_attempted", result)
        self.assertIn("production", result["accuracy_testing_warning"])
        path_exists.assert_not_called()
        load_pending.assert_not_called()
        queue_analysis.assert_not_called()

    def test_disabled_local_export_stops_before_filesystem_access(self):
        with patch.dict(
            os.environ,
            {"APP_ENV": "testing", "ACCURACY_TESTING_AUTO_EXPORT": "0"},
            clear=False,
        ), patch.object(app_module.os.path, "exists") as path_exists:
            result = app_module.update_accuracy_testing_workbook(sample_analysis())

        self.assertFalse(result["accuracy_testing_saved"])
        self.assertNotIn("accuracy_testing_attempted", result)
        self.assertIn("disabled", result["accuracy_testing_warning"])
        path_exists.assert_not_called()

    def test_production_health_hides_workbook_diagnostics(self):
        client = app_module.app.test_client()
        with patch.dict(os.environ, {"APP_ENV": "production"}, clear=False):
            response = client.get("/health")

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertNotIn("accuracy_testing_workbook", payload)
        self.assertNotIn("accuracy_testing_workbook_exists", payload)
        self.assertNotIn("accuracy_testing_pending_count", payload)


if __name__ == "__main__":
    unittest.main()
